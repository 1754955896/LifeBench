"""Read common source containers before semantic LLM normalization."""

import csv
import json
import os
from pathlib import Path
from typing import Any, List, Optional


def _as_records(value: Any) -> List[Any]:
    return list(value) if isinstance(value, list) else [value]


def load_source_records(source: Any, input_format: str = "auto") -> List[Any]:
    if not isinstance(source, (str, os.PathLike)) or not os.path.exists(os.fspath(source)):
        return _as_records(source)

    path = Path(source)
    fmt = input_format.lower()
    if fmt == "auto":
        fmt = path.suffix.lower().lstrip(".") or "text"
    if fmt == "json":
        with path.open("r", encoding="utf-8") as stream:
            return _as_records(json.load(stream))
    if fmt in {"jsonl", "ndjson"}:
        with path.open("r", encoding="utf-8") as stream:
            return [json.loads(line) for line in stream if line.strip()]
    if fmt in {"csv", "tsv"}:
        delimiter = "\t" if fmt == "tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream, delimiter=delimiter))
    if fmt in {"txt", "text", "md"}:
        return [path.read_text(encoding="utf-8")]
    if fmt in {"xlsx", "xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("读取 Excel 需要安装 openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]
    raise ValueError("不支持的输入格式: %s" % fmt)
